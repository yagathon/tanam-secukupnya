"""A year of farm days for the sandbox: make an editable Excel scenario, then play it into a running sandbox.

    uv run --project backend python scripts/scenario.py make                 # → data/scenario/year.xlsx
    uv run --project backend python scripts/scenario.py export               # xlsx → year.json (for the UI bot)
    uv run --project backend python scripts/scenario.py play --reset         # clear the sandbox, load the setup, play all
    uv run --project backend python scripts/scenario.py play --days 30       # continue from the sandbox's today

The orders come from the example farm's own demand (the founder's procurement numbers: Resto WKS' week,
the payday month, the year index), with day-to-day noise, a slowly growing retail pre-order list and a few
dated stories: rain losses in the wet season, Ramadan and Lebaran (announced as signals), and one drop
nobody announced. Sow rows say "plan" = sow what the app proposes that morning; type a number to override.
Harvests are not in the sheet: a day without one counts the app's planned send as harvested.
"""

from __future__ import annotations

import argparse
import json
import math
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app import sandbox as sb  # noqa: E402
from app.demand import season_factor  # noqa: E402

OUT = ROOT / "data" / "scenario"
XLSX, JSON = OUT / "year.xlsx", OUT / "year.json"
START = date(2026, 10, 17)           # the founder's real beds are aged from this day
DAYS = 365
COLS = ["date", "kind", "buyer", "line", "kg", "for_day", "seedlings", "pct", "who", "effect_pct",
        "from", "to", "cause", "note"]
KINDS = {"order", "sow", "loss", "signal", "postpone"}

# Dated stories, (from, to) inclusive. Ramadan 1448 ≈ 8 Feb – 9 Mar 2027, Idul Fitri ≈ 10 Mar 2027.
RAMADAN = (date(2027, 2, 8), date(2027, 3, 9))
LEBARAN_CLOSED = (date(2027, 3, 8), date(2027, 3, 14))
LEBARAN_RETAIL = (date(2027, 3, 6), date(2027, 3, 7))
SILENT_DROP = (date(2027, 5, 10), date(2027, 5, 23))      # Resto WKS orders −35%, nobody says why


def _in(d: date, span: tuple[date, date]) -> bool:
    return span[0] <= d <= span[1]


def truth(b: dict, d: date, season: dict, year_frac: float) -> tuple[float, str]:
    """The buyer's expected kg on day d before noise, and the story behind it (empty = an ordinary day)."""
    kg = sb._week(b)[d.weekday()] * season_factor(season, d)
    if b["name"] == "Retails":
        kg *= 1 + 0.8 * year_frac                         # the pre-order list grows: 5 → 9 kg a week
        if _in(d, LEBARAN_RETAIL):
            return kg * 1.5, "Lebaran pre-orders +50%"
        return kg, ""
    if _in(d, LEBARAN_CLOSED):
        return 0.0, "closed for Lebaran"
    if _in(d, RAMADAN):
        return kg * 0.8, "Ramadan −20%"
    if _in(d, SILENT_DROP):
        return kg * 0.65, "unannounced drop −35%"
    return kg, ""


def make(path: Path = XLSX, seed: int = 17, start: date = START, days: int = DAYS) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rng = random.Random(seed)
    setup = sb.example_setup(start)
    season = {"month_pct": setup["month_pct"], "year_index": setup["year_index"]}
    line = setup["crop"]["lines"][0]["id"]
    sow_days = set(setup["crop"]["sow_weekdays"])
    rows: list[dict] = []
    truth_rows: list[list] = []
    rain_days = {start + timedelta(days=i) for i in range(days)
                 if (start + timedelta(days=i)).month in (11, 12, 1, 2, 3) and rng.random() < 0.06}

    def add(d: date, kind: str, **kw):
        rows.append({"date": d, "kind": kind, "line": line, **kw})

    for i in range(days):
        d = start + timedelta(days=i)
        nxt = d + timedelta(days=1)
        if i == days - 1:
            break
        # orders: Resto WKS orders the day before; Retails pre-orders Thursday for Saturday and Sunday
        for b in setup["buyers"]:
            if b["name"] == "Retails":
                if d.weekday() != 3:
                    continue
                targets = [d + timedelta(days=2), d + timedelta(days=3)]
            else:
                targets = [nxt]
            for day in targets:
                if day >= start + timedelta(days=days):
                    continue
                kg, story = truth(b, day, season, (day - start).days / days)
                sigma = 0.3 if b["name"] == "Retails" else 0.15
                ordered = 0.0 if kg == 0 else round(kg * math.exp(rng.gauss(0, sigma)) * 2) / 2
                if b["name"] != "Retails" and kg and rng.random() < 0.02:
                    ordered, story = 0.0, "skipped this day"
                if kg or story:
                    add(d, "order", buyer=b["name"], kg=ordered, for_day=day, note=story)
                truth_rows.append([day, day.strftime("%a"), b["name"], round(kg, 2), ordered, story])
        # the wet season: heavy rain takes 10–25% of the plants, and sometimes stops a sowing
        if d in rain_days:
            add(d, "loss", pct=rng.choice([10, 15, 20, 25]), cause="hujan deras", note="wet season")
        if d.weekday() in sow_days:
            if d in rain_days and rng.random() < 0.5:
                add(d, "postpone", cause="hujan, bibit belum bisa pindah", note="wet season")
            else:
                add(d, "sow", seedlings="plan")
        if d == date(2027, 6, 14):
            add(d, "loss", pct=15, cause="ulat", note="caterpillars")
        # the stories people do tell: Ramadan and Lebaran are announced a week or so ahead
        if d == RAMADAN[0] - timedelta(days=7):
            add(d, "signal", line="", who="Resto WKS", effect_pct=-20, **{"from": RAMADAN[0], "to": LEBARAN_CLOSED[0] - timedelta(days=1)},
                cause="Ramadan: lunch trade drops", note="announced")
        if d == LEBARAN_CLOSED[0] - timedelta(days=10):
            add(d, "signal", line="", who="Resto WKS", effect_pct=-100, **{"from": LEBARAN_CLOSED[0], "to": LEBARAN_CLOSED[1]},
                cause="tutup Lebaran", note="announced")
        if d == LEBARAN_RETAIL[0] - timedelta(days=7):
            add(d, "signal", line="", who="Retails", effect_pct=50, **{"from": LEBARAN_RETAIL[0], "to": LEBARAN_RETAIL[1]},
                cause="pre-order Lebaran", note="announced")

    wb = Workbook()
    ws = wb.active
    ws.title = "Days"
    ws.append(COLS)
    for r in rows:
        ws.append([r.get(c) for c in COLS])
    for row in ws.iter_rows(min_row=2):
        for c in (row[0], row[5], row[10], row[11]):
            if c.value:
                c.number_format = "yyyy-mm-dd ddd"
    for c, w in zip("ABCDEFGHIJKLMN", (16, 9, 11, 6, 6, 16, 10, 5, 11, 10, 16, 16, 30, 26)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A2"
    tr = wb.create_sheet("Truth")
    tr.append(["day", "weekday", "buyer", "expected_kg", "ordered_kg", "story"])
    for r in sorted(truth_rows, key=lambda r: (r[0], r[2])):
        tr.append(r)
        tr.cell(tr.max_row, 1).number_format = "yyyy-mm-dd"
    st = wb.create_sheet("Setup")
    st["A1"] = json.dumps(setup)
    st["A2"] = "The sandbox setup the play starts from (JSON). Edit it in the app instead, then paste its JSON here."
    rd = wb.create_sheet("Read me", 0)
    for line_ in __doc__.strip().splitlines() + ["", "Days columns:",
                 "order: buyer, kg (0 = no order this time), for_day",
                 "sow: seedlings = plan (follow the app) or a number",
                 "loss: pct, cause (all beds of the line)",
                 "signal: who (everyone / businesses / households / a buyer name), effect_pct, from, to, cause",
                 "postpone: cause (today's sowing does not happen)",
                 "note: for you only; the play ignores it. Delete or add rows freely, keep the date order."]:
        rd.append([line_])
    rd["A1"].font = Font(bold=True)
    rd.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"rows": len(rows), "path": str(path)}


def _iso(v) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return date.fromisoformat(str(v).strip()[:10]).isoformat()


def load(path: Path = XLSX) -> dict:
    """The scenario as plain data: {"setup": {...}, "rows": [{date, kind, ...}]} in date order."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    setup = json.loads(wb["Setup"]["A1"].value)
    rows = []
    for n, vals in enumerate(wb["Days"].iter_rows(min_row=2, values_only=True), start=2):
        r = dict(zip(COLS, vals))
        if not r["date"] and not r["kind"]:
            continue
        kind = str(r["kind"] or "").strip().lower()
        if kind not in KINDS:
            raise ValueError(f"Days row {n}: unknown kind {r['kind']!r}")
        for k in ("date", "for_day", "from", "to"):
            r[k] = _iso(r[k])
        r["kind"], r["row"] = kind, n
        rows.append({k: v for k, v in r.items() if v not in (None, "")})
    rows.sort(key=lambda r: r["date"])                      # stable: same-day rows keep the sheet's order
    return {"setup": setup, "rows": rows}


def form(r: dict, ids: dict[str, str], sow_plan: dict[str, tuple[str, int]], today: str) -> dict | None:
    """One sheet row → the sandbox's form body (what a person types), or None to skip it."""
    k = r["kind"]
    if k == "order":
        return {"kind": "order", "buyer": ids[r["buyer"]], "line": r["line"], "kg": r["kg"], "day": r["for_day"]}
    if k == "sow":
        n = r.get("seedlings", "plan")
        if str(n).strip().lower() == "plan":
            day, n = sow_plan.get(r["line"], ("", 0))
            if day != today or n <= 0:
                return None
        return {"kind": "sow", "line": r["line"], "seedlings": int(n)}
    if k == "loss":
        return {"kind": "loss", "line": r["line"], "batch": "", "pct": r["pct"], "cause": r.get("cause", "")}
    if k == "postpone":
        return {"kind": "postpone", "line": r["line"], "reason": r.get("cause", "")}
    who = str(r.get("who", "everyone"))
    scope = {"everyone": "all", "businesses": "channel:B2B", "households": "channel:B2C"}.get(who.lower()) \
        or f"buyer:{ids[who]}"
    return {"kind": "signal", "scope": scope, "line": r.get("line", ""), "effect_pct": r["effect_pct"],
            "date_from": r["from"], "date_to": r.get("to", ""), "cause": r.get("cause", "")}


def play(api, scenario: dict, reset: bool = False, days: int | None = None, say=print) -> dict:
    """Play the rows into the sandbox day by day from its current today. api: get(path) / post(path, body) /
    put(path, body) returning parsed JSON, raising on an error status (see Http)."""
    if reset:
        api.post("/clear", {"setup": True})
        api.put("/setup", scenario["setup"])
    s = api.get("/setup")
    if not s["saved"]:
        raise SystemExit("the sandbox has no setup: run with --reset")
    ids = {b["name"]: b["id"] for b in s["setup"]["buyers"]}
    by_day: dict[str, list[dict]] = {}
    for r in scenario["rows"]:
        by_day.setdefault(r["date"], []).append(r)
    last = max(by_day) if by_day else None
    played = skipped = 0
    n = 0
    while True:
        t = api.get("/today")
        today, seq = t["today"], t["seq"]
        if last is None or today > last or (days is not None and n >= days):
            break
        plan = {l["line"]: (l["day"], l["seedlings"]) for l in api.get("/sow")["lines"]}
        for r in by_day.get(today, []):
            body = form(r, ids, plan, today)
            if body is None:
                skipped += 1
                continue
            try:
                api.post("/events", {"form": body, "seq": seq})
                seq += 1
                played += 1
            except HttpError as e:
                skipped += 1
                say(f"  {today} row {r['row']} {r['kind']}: skipped ({e})")
        api.post("/day/next", {"from_day": today, "days": 1})
        n += 1
        if n % 30 == 0:
            say(f"day {n}: {today}")
    return {"days": n, "events": played, "skipped": skipped, "today": api.get("/today")["today"]}


class HttpError(Exception):
    pass


class Http:
    """The running sandbox API (local: `just api`, AUTH=off)."""

    def __init__(self, base: str, client=None):
        import requests

        self.base, self.c = base.rstrip("/") + "/api/sandbox", client or requests.Session()

    def _do(self, method: str, path: str, body=None):
        r = self.c.request(method, self.base + path, json=body)
        if r.status_code >= 400:
            try:
                msg = r.json().get("detail", r.text)
            except ValueError:
                msg = r.text
            raise HttpError(f"{r.status_code} {msg}")
        return r.json()

    def get(self, path):
        return self._do("GET", path)

    def post(self, path, body):
        return self._do("POST", path, body)

    def put(self, path, body):
        return self._do("PUT", path, body)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("make")
    m.add_argument("--seed", type=int, default=17)
    m.add_argument("--out", type=Path, default=XLSX)
    e = sub.add_parser("export")
    e.add_argument("--xlsx", type=Path, default=XLSX)
    e.add_argument("--out", type=Path, default=JSON)
    pl = sub.add_parser("play")
    pl.add_argument("--xlsx", type=Path, default=XLSX)
    pl.add_argument("--api", default="http://localhost:8100")
    pl.add_argument("--reset", action="store_true", help="clear the sandbox (setup too) and load the sheet's setup")
    pl.add_argument("--days", type=int, help="play this many days, then stop (default: to the end)")
    a = p.parse_args()
    if a.cmd == "make":
        print(make(a.out, a.seed))
    elif a.cmd == "export":
        a.out.write_text(json.dumps(load(a.xlsx), indent=1))
        print(f"wrote {a.out}")
    else:
        print(play(Http(a.api), load(a.xlsx), a.reset, a.days))


if __name__ == "__main__":
    main()
